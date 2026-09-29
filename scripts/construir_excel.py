"""
construir_excel.py | Construye excel/analisis_ventas.xlsx a partir de la base limpia
=====================================================================================
Requisito previo: haber ejecutado verificar_con_sqlite.py (o exportar v_ventas desde MySQL
a excel/base_pedidos.csv con las mismas columnas).

El libro usa FÓRMULAS (SUMIFS, COUNTIFS, INDEX...) sobre la hoja Base_Pedidos, de modo que
si cambias los datos, los resúmenes se recalculan. Cada resumen se contrasta con el SQL.

Uso:  python construir_excel.py
"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
base = pd.read_csv(RAIZ / "excel" / "base_pedidos.csv", parse_dates=["fecha_pedido"])
res = {q: pd.read_csv(RAIZ / "resultados" / f"{q}.csv") for q in ["Q02", "Q07", "Q08", "Q09"]}
N = len(base)
ULT = N + 1                                     # última fila de datos en Base_Pedidos

FUENTE = "Arial"
F_NORMAL, F_BOLD = Font(name=FUENTE, size=10), Font(name=FUENTE, size=10, bold=True)
F_HEAD = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FUENTE, size=10, color="0000FF")          # azul = valor escrito a mano / viene del SQL
F_TITULO = Font(name=FUENTE, size=14, bold=True, color="1F3864")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_TOTAL = PatternFill("solid", fgColor="D9E1F2")
BORDE = Border(bottom=Side(style="thin", color="BFBFBF"))
COP, PCT, ENT = '#,##0', '0.0%', '#,##0'

wb = Workbook()


def encabezado(ws, fila, textos, ancho=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(row=fila, column=j, value=t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 32
    if ancho:
        for j, w in enumerate(ancho, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


# ------------------------------------------------------------------ LEEME
ws = wb.active
ws.title = "LEEME"
ws["A1"], ws["A1"].font = "Análisis de ventas y rentabilidad - Tienda en línea", F_TITULO
lineas = [
    ("Qué es", "Resumen en Excel del análisis hecho en SQL (ver carpeta sql/). Datos SINTÉTICOS generados con scripts/generar_datos.py."),
    ("Periodo", "1-jul-2024 a 31-dic-2025. Moneda: pesos colombianos (COP)."),
    ("", ""),
    ("HOJAS", ""),
    ("Base_Pedidos", "41.407 pedidos ENTREGADOS ya limpios (salida de la vista v_ventas). Punto de partida de todo el libro."),
    ("Resumen_Categorias", "Ingresos, participación, devoluciones y margen por categoría (SUMIFS / COUNTIFS)."),
    ("Ingresos_Mensual", "Ingresos por mes, acumulado y variación mensual (equivale a Q03 y Q04 del SQL)."),
    ("Pareto", "Ranking de productos y % acumulado de ingresos (equivale a Q08 del SQL)."),
    ("Verificacion_SQL", "Compara los totales de Excel contra los resultados de SQL. Todo debe decir OK."),
    ("", ""),
    ("DEFINICIONES", ""),
    ("Ingreso", "Valor efectivamente pagado (tabla pagos) de pedidos con estado 'entregado'."),
    ("Ingreso neto", "Ingreso - reembolsos por devoluciones."),
    ("Utilidad", "Ingreso neto - costo (cantidad x costo unitario). Supuesto: lo devuelto no se revende."),
    ("Tasa de devolución", "Pedidos devueltos / pedidos entregados."),
    ("", ""),
    ("CÓMO CREAR TU TABLA DINÁMICA", "(ejercicio recomendado: no está pre-creada para que practiques)"),
    ("Paso 1", "Ve a Base_Pedidos, haz clic en cualquier celda de los datos."),
    ("Paso 2", "Menú Insertar > Tabla dinámica > Desde tabla o rango > Nueva hoja de cálculo."),
    ("Paso 3", "Arrastra 'categoria' a Filas y 'ingreso' a Valores (Suma). Arrastra 'mes' a Columnas."),
    ("Paso 4", "Añade 'devuelto' a Valores con 'Promedio': eso es la tasa de devolución."),
    ("Paso 5", "Insertar > Gráfico dinámico. Compara tu resultado con la hoja Resumen_Categorias."),
    ("", ""),
    ("Colores", "Texto azul = valor escrito a mano o copiado del SQL. Texto negro = fórmula."),
]
for i, (a, b) in enumerate(lineas, 3):
    ws.cell(row=i, column=1, value=a).font = F_BOLD
    ws.cell(row=i, column=2, value=b).font = F_NORMAL
    ws.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 110

# ------------------------------------------------------------------ BASE_PEDIDOS
wb_base = wb.create_sheet("Base_Pedidos")
cols = ["pedido_id", "cliente_id", "producto_id", "nombre_producto", "categoria", "fecha_pedido",
        "cantidad", "metodo_pago", "ingreso", "costo", "reembolso", "devuelto", "ingreso_neto", "utilidad", "mes"]
encabezado(wb_base, 1, cols, [11, 11, 12, 16, 14, 13, 10, 20, 15, 15, 14, 10, 15, 15, 10])
for i, r in enumerate(base.itertuples(index=False), 2):
    wb_base.append([r.pedido_id, r.cliente_id, r.producto_id, r.nombre_producto, r.categoria,
                    r.fecha_pedido.to_pydatetime().date(), r.cantidad, r.metodo_pago,
                    r.ingreso, r.costo, r.reembolso, r.devuelto,
                    f"=I{i}-K{i}", f"=I{i}-K{i}-J{i}", f'=YEAR(F{i})&"-"&TEXT(MONTH(F{i}),"00")'])
for row in wb_base.iter_rows(min_row=2, max_row=ULT):
    for c in row:
        c.font = F_NORMAL
    row[5].number_format = "yyyy-mm-dd"
    for k in (8, 9, 10, 12, 13):
        row[k].number_format = COP
wb_base.freeze_panes = "A2"
wb_base.auto_filter.ref = f"A1:O{ULT}"


def rng(col):
    return f"Base_Pedidos!${col}$2:${col}${ULT}"


# ------------------------------------------------------------------ RESUMEN_CATEGORIAS
wr = wb.create_sheet("Resumen_Categorias")
wr["A1"], wr["A1"].font = "Resumen por categoría (equivale a Q07 y Q09 del SQL)", F_TITULO
encabezado(wr, 3, ["Categoría", "Pedidos entregados", "Ingresos (COP)", "Participación", "Reembolsos (COP)",
                   "Pedidos devueltos", "Tasa de devolución", "Utilidad (COP)", "Margen s/ ingreso neto"],
           [18, 14, 18, 14, 18, 14, 14, 18, 16])
cats = res["Q07"]["categoria"].tolist()                        # orden = por ingresos desc (Q07)
f0, fN = 4, 3 + len(cats)
tot = fN + 1
for i, cat in enumerate(cats, f0):
    wr.cell(row=i, column=1, value=cat)
    wr.cell(row=i, column=2, value=f"=COUNTIFS({rng('E')},A{i})")
    wr.cell(row=i, column=3, value=f"=SUMIFS({rng('I')},{rng('E')},A{i})")
    wr.cell(row=i, column=4, value=f"=C{i}/C${tot}")
    wr.cell(row=i, column=5, value=f"=SUMIFS({rng('K')},{rng('E')},A{i})")
    wr.cell(row=i, column=6, value=f"=SUMIFS({rng('L')},{rng('E')},A{i})")
    wr.cell(row=i, column=7, value=f"=F{i}/B{i}")
    wr.cell(row=i, column=8, value=f"=SUMIFS({rng('N')},{rng('E')},A{i})")
    wr.cell(row=i, column=9, value=f"=H{i}/(C{i}-E{i})")
wr.cell(row=tot, column=1, value="TOTAL")
for col in "BCEFH":
    wr[f"{col}{tot}"] = f"=SUM({col}{f0}:{col}{fN})"
wr[f"D{tot}"] = f"=C{tot}/C{tot}"
wr[f"G{tot}"] = f"=F{tot}/B{tot}"
wr[f"I{tot}"] = f"=H{tot}/(C{tot}-E{tot})"
for r in range(f0, tot + 1):
    for c in range(1, 10):
        cell = wr.cell(row=r, column=c)
        cell.font = F_BOLD if r == tot else F_NORMAL
        cell.border = BORDE
        if r == tot:
            cell.fill = FILL_TOTAL
        cell.number_format = {2: ENT, 3: COP, 4: PCT, 5: COP, 6: ENT, 7: '0.00%', 8: COP, 9: PCT}.get(c, "General")
ch = BarChart()
ch.type, ch.title, ch.style = "col", "Tasa de devolución por categoría", 10
ch.add_data(Reference(wr, min_col=7, min_row=3, max_row=fN), titles_from_data=True)
ch.set_categories(Reference(wr, min_col=1, min_row=f0, max_row=fN))
ch.y_axis.title, ch.y_axis.numFmt, ch.legend = "Tasa", "0%", None
ch.x_axis.delete = ch.y_axis.delete = False
ch.height, ch.width = 8, 16
wr.add_chart(ch, "A14")
ch2 = BarChart()
ch2.type, ch2.title, ch2.style = "col", "Ingresos por categoría (COP)", 10
ch2.add_data(Reference(wr, min_col=3, min_row=3, max_row=fN), titles_from_data=True)
ch2.set_categories(Reference(wr, min_col=1, min_row=f0, max_row=fN))
ch2.legend = None
ch2.x_axis.delete = ch2.y_axis.delete = False
ch2.height, ch2.width = 8, 16
wr.add_chart(ch2, "F14")

# ------------------------------------------------------------------ INGRESOS_MENSUAL
wm = wb.create_sheet("Ingresos_Mensual")
wm["A1"], wm["A1"].font = "Ingresos por mes (equivale a Q03 y Q04 del SQL)", F_TITULO
encabezado(wm, 3, ["Mes", "Ingresos (COP)", "Acumulado (COP)", "Variación vs mes anterior"], [12, 18, 18, 16])
meses = sorted(base["fecha_pedido"].dt.strftime("%Y-%m").unique())
m0, mN = 4, 3 + len(meses)
for i, mes in enumerate(meses, m0):
    wm.cell(row=i, column=1, value=mes)
    wm.cell(row=i, column=2, value=f"=SUMIFS({rng('I')},{rng('O')},A{i})")
    wm.cell(row=i, column=3, value=f"=SUM(B${m0}:B{i})")
    if i > m0:
        wm.cell(row=i, column=4, value=f"=B{i}/B{i-1}-1")
    for c, nf in zip(range(1, 5), ["General", COP, COP, PCT]):
        wm.cell(row=i, column=c).font = F_NORMAL
        wm.cell(row=i, column=c).number_format = nf
        wm.cell(row=i, column=c).border = BORDE
lc = LineChart()
lc.title, lc.style, lc.legend = "Ingresos mensuales (COP)", 12, None
lc.add_data(Reference(wm, min_col=2, min_row=3, max_row=mN), titles_from_data=True)
lc.set_categories(Reference(wm, min_col=1, min_row=m0, max_row=mN))
lc.x_axis.delete = lc.y_axis.delete = False
lc.height, lc.width = 8, 18
wm.add_chart(lc, "F3")

# ------------------------------------------------------------------ PARETO
wp = wb.create_sheet("Pareto")
wp["A1"], wp["A1"].font = "Análisis de Pareto de productos (equivale a Q08 del SQL)", F_TITULO
encabezado(wp, 3, ["N.º", "producto_id", "Producto", "Categoría", "Ingresos (COP)", "% acumulado de ingresos"],
           [8, 12, 14, 14, 18, 16])
orden = base.groupby(["producto_id", "nombre_producto", "categoria"]).ingreso.sum().reset_index() \
            .sort_values(["ingreso", "producto_id"], ascending=[False, True]).reset_index(drop=True)
p0, pN = 4, 3 + len(orden)
for i, r in enumerate(orden.itertuples(index=False), p0):
    wp.cell(row=i, column=1, value=i - p0 + 1)
    wp.cell(row=i, column=2, value=int(r.producto_id))
    wp.cell(row=i, column=3, value=r.nombre_producto)
    wp.cell(row=i, column=4, value=r.categoria)
    wp.cell(row=i, column=5, value=f"=SUMIFS({rng('I')},{rng('C')},B{i})")
    wp.cell(row=i, column=6, value=f"=SUM(E${p0}:E{i})/SUM(E${p0}:E${pN})")
    for c, nf in zip(range(1, 7), ["General", "General", "General", "General", COP, PCT]):
        wp.cell(row=i, column=c).font = F_NORMAL
        wp.cell(row=i, column=c).number_format = nf
wp["H3"], wp["H3"].font = "RESULTADOS CLAVE", F_BOLD
claves = [
    ("Productos en el catálogo (tabla productos, valor del SQL)", 500, "0"),
    ("Productos necesarios para el 70 % de ingresos", f'=COUNTIF(F{p0}:F{pN},"<0.7")+1', "0"),
    ("... como % del catálogo", "=I5/I4", PCT),
    ("% de ingresos del 20 % de productos con más ventas", f"=INDEX(F{p0}:F{pN},ROUND(0.2*I4,0))", PCT),
]
for i, (t, f, nf) in enumerate(claves, 4):
    wp.cell(row=i, column=8, value=t).font = F_NORMAL
    c = wp.cell(row=i, column=9, value=f)
    c.font, c.number_format = (F_INPUT if i == 4 else F_BOLD), nf
wp.column_dimensions["H"].width = 58
wp.column_dimensions["I"].width = 12
wp["H9"] = "Nota: el orden del ranking se fijó al construir el libro; si cambias los datos, reordena por Ingresos."
wp["H9"].font = Font(name=FUENTE, size=9, italic=True)
pc = LineChart()
pc.title, pc.style, pc.legend = "Curva de Pareto: % acumulado de ingresos", 12, None
pc.add_data(Reference(wp, min_col=6, min_row=3, max_row=pN), titles_from_data=True)
pc.x_axis.title, pc.y_axis.title, pc.y_axis.numFmt = "Productos (de mayor a menor venta)", "% acumulado", "0%"
pc.x_axis.delete = pc.y_axis.delete = False
pc.x_axis.tickLblSkip = 50
pc.height, pc.width = 9, 18
wp.add_chart(pc, "H11")

# ------------------------------------------------------------------ VERIFICACION_SQL
wv = wb.create_sheet("Verificacion_SQL")
wv["A1"], wv["A1"].font = "Verificación: Excel vs. SQL", F_TITULO
wv["A2"] = "Valor SQL (azul) = resultado de las consultas Q02, Q08 y Q09 (carpeta resultados/). Valor Excel = fórmula sobre Base_Pedidos."
wv["A2"].font = Font(name=FUENTE, size=9, italic=True)
encabezado(wv, 4, ["Métrica", "Valor SQL", "Valor Excel", "Diferencia", "¿Coincide?"], [44, 20, 20, 14, 14])
q2, q8, q9 = res["Q02"].iloc[0], res["Q08"].iloc[0], res["Q09"].set_index("categoria")
filas = [
    ("Pedidos entregados", int(q2.pedidos_entregados), f"=COUNTA({rng('A')})", ENT, 0.5),
    ("Ingresos (COP)", float(q2.ingresos), f"=SUM({rng('I')})", COP, 0.5),
    ("Reembolsos (COP)", float(q2.reembolsos), f"=SUM({rng('K')})", COP, 0.5),
    ("Ingreso neto (COP)", float(q2.ingreso_neto), f"=SUM({rng('M')})", COP, 0.5),
    ("Productos para el 70 % de ingresos", int(q8.productos_para_70pct_ingresos), "=Pareto!I5", "0", 0.5),
    ("Tasa de devolución Ropa (%)", float(q9.loc["Ropa", "tasa_devolucion_pct"]),
     f"=ROUND(Resumen_Categorias!G{cats.index('Ropa') + f0}*100,2)", "0.00", 0.01),
    ("Tasa de devolución Electrónica (%)", float(q9.loc["Electrónica", "tasa_devolucion_pct"]),
     f"=ROUND(Resumen_Categorias!G{cats.index('Electrónica') + f0}*100,2)", "0.00", 0.01),
]
for i, (t, sqlv, f, nf, tol) in enumerate(filas, 5):
    wv.cell(row=i, column=1, value=t).font = F_NORMAL
    c = wv.cell(row=i, column=2, value=sqlv)
    c.font, c.number_format = F_INPUT, nf
    c = wv.cell(row=i, column=3, value=f)
    c.font, c.number_format = F_NORMAL, nf
    c = wv.cell(row=i, column=4, value=f"=C{i}-B{i}")
    c.font, c.number_format = F_NORMAL, "#,##0.00"
    c = wv.cell(row=i, column=5, value=f'=IF(ABS(D{i})<={tol},"OK","REVISAR")')
    c.font, c.alignment = F_BOLD, Alignment(horizontal="center")

for hoja in wb.worksheets:
    hoja.sheet_view.showGridLines = hoja.title in ("Base_Pedidos",)
out = RAIZ / "excel" / "analisis_ventas.xlsx"
wb.save(out)
print("Guardado:", out)
